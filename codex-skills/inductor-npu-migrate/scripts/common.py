"""Shared deterministic file discovery and workbook conventions."""
from pathlib import Path
import ast
import hashlib
import json
import tokenize


def read_source(path):
    with tokenize.open(path) as stream:
        return stream.read()


def test_files(root):
    return sorted(p for p in Path(root).rglob('test_*.py') if p.is_file())


def cases(path):
    """Static definitions; never import torch or execute parameterization."""
    result = []
    def visit(node, classes=()):
        if isinstance(node, ast.ClassDef):
            classes += (node.name,)
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            if node.name.startswith('test_'):
                result.append((node.name, '.'.join(classes)))
            # A function-local test/helper is not a separately collected case.
            return
        for child in ast.iter_child_nodes(node):
            visit(child, classes)
    visit(ast.parse(read_source(path), filename=str(path)))
    return result


def digest(data):
    return hashlib.sha256(data).hexdigest()


def safe_path(root, relative):
    root = Path(root).resolve()
    path = (root / relative).resolve()
    if path == root or not path.is_relative_to(root):
        raise ValueError(f'Path escapes root: {relative}')
    return path


def write_json(path, obj):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(obj, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def style_sheet(sheet):
    from openpyxl.styles import Alignment, Font, PatternFill
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
    for i, width in enumerate((42, 60, 42, 18, 100), 1):
        sheet.column_dimensions[sheet.cell(1, i).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
