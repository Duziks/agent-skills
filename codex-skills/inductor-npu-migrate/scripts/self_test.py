#!/usr/bin/env python3
"""Offline fixture integration checks. Does not run pytest or import torch."""
import argparse
from collections import Counter
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from openpyxl import load_workbook
from common import digest, write_json
from fill_test_results import fill
from summarize_pytest_logs import parse_log

S = Path(__file__).resolve().parent


def run(script, *args, ok=True):
    environment = dict(os.environ, PYTHONIOENCODING='utf-8')
    result = subprocess.run([sys.executable, str(S / script), *map(str, args)], capture_output=True,
                            text=True, encoding='utf-8', env=environment)
    if (result.returncode == 0) != ok:
        raise AssertionError(f'{script} exit {result.returncode}\n{result.stdout}\n{result.stderr}')
    return result


def check(root):
    root.mkdir(parents=True, exist_ok=True)
    repo = root / 'fixture_repo'
    repo.mkdir()
    def git(*args):
        return subprocess.run(['git', '-C', str(repo), *args], capture_output=True, text=True, check=True).stdout.strip()
    git('init')
    git('config', 'user.email', 'fixture@example.invalid')
    git('config', 'user.name', 'Offline fixture')
    git('remote', 'add', 'origin', 'https://github.com/pytorch/pytorch.git')
    tests = repo / 'test' / 'inductor'
    tests.mkdir(parents=True)
    source = '''#!/usr/bin/env python3
# coding: utf-8
# Fixture only
"""Fixture module docstring."""
from __future__ import annotations

class TestBase:
    def test_mix(self):
        return "cuda"
    def test_skip(self): pass
    def test_pass(self): pass
    def test_missing(self): pass
    def test_error(self): pass
    def test_xfail(self): pass
    def test_xpass(self): pass
    def test_add(self): pass
    def helper(self):
        def test_inner_not_collected(): pass

class Another:
    def test_pass(self): pass
'''
    (tests / 'test_a.py').write_text(source, encoding='utf-8')
    nested = tests / 'nested'
    nested.mkdir()
    (nested / 'test_a.py').write_text('def test_pass(): pass\n', encoding='utf-8')
    for name in ('test_collect', 'test_interrupt', 'test_exit'):
        (tests / f'{name}.py').write_text('def test_pass(): pass\n', encoding='utf-8')
    (tests / 'helper.py').write_text('def test_not_a_test_file(): pass\n', encoding='utf-8')
    (tests / 'fixture.txt').write_text('unchanged resource', encoding='utf-8')
    (repo / 'test' / 'common_helper.py').write_text('# retained helper outside inductor\n', encoding='utf-8')
    git('add', '.')
    git('commit', '-m', 'offline fixture')
    git('tag', 'v0.0.1')
    commit = git('rev-parse', 'HEAD')
    (tests / 'test_a.py').write_text('# dirty content must not enter export\n', encoding='utf-8')
    output = root / 'fixture_run'
    run('prepare_source.py', '--tag', 'v0.0.1', '--source-repo', repo, '--output', output)
    manifest = json.loads((output / 'source_manifest.json').read_text(encoding='utf-8'))
    assert manifest['commit'] == commit
    assert manifest['test_files'] == ['nested/test_a.py', 'test_a.py', 'test_collect.py', 'test_exit.py', 'test_interrupt.py']
    target = output / 'adapted' / 'test' / 'inductor'
    assert (target / 'test_a.py').read_text(encoding='utf-8') == source
    assert (target / 'fixture.txt').read_text(encoding='utf-8') == 'unchanged resource'
    assert (target.parent / 'common_helper.py').exists()
    inventory = output / '待测用例.xlsx'
    book = load_workbook(inventory)
    sheet = book.active
    base_cells = {(r, c): sheet.cell(r, c).value for r in range(1, sheet.max_row + 1) for c in range(1, 4)}
    assert sheet.max_row == 14  # 13 definitions + header
    assert not any('test_inner' in str(row[1].value) for row in sheet.iter_rows(min_row=2))
    sheet['A2'].number_format = '@'
    book.create_sheet('preserved')['B2'] = 'keep'
    book.save(inventory)
    book.close()
    entries = []
    for name in manifest['test_files']:
        entry = dict(path=name, sha256=digest((target / name).read_bytes()), status='unchanged',
                     reason='Offline fixture classification', intents=[dict(symbol=name, status='already_satisfied', evidence='fixture pure logic')])
        if name == 'test_a.py':
            entry.update(status='preliminary', initialization='import torch\nimport torch_npu\nimport torch_npu._inductor',
                         operations=[dict(old='        return "cuda"', new='        return "npu"')])
            entry['intents'][0]['status'] = 'applied_semantic'
        entries.append(entry)
    plan, report = output / 'adaptation_plan.json', output / 'adaptation_report.json'
    write_json(plan, dict(files=entries))
    run('apply_adaptation.py', '--root', target, '--plan', plan, '--report', report)
    adapted_bytes = (target / 'test_a.py').read_bytes()
    assert adapted_bytes.index(b'from __future__') < adapted_bytes.index(b'import torch_npu')
    assert b'return "npu"' in adapted_bytes
    assert (output / 'original' / 'test' / 'inductor' / 'test_a.py').read_text(encoding='utf-8') == source
    run('apply_adaptation.py', '--root', target, '--plan', plan, '--report', report)
    assert adapted_bytes == (target / 'test_a.py').read_bytes()
    # Hash mismatch must fail without changing a source.
    altered_plan = output / 'bad_plan.json'
    entries[0]['sha256'] = '0' * 64
    write_json(altered_plan, dict(files=entries))
    run('apply_adaptation.py', '--root', target, '--plan', altered_plan, '--report', output / 'bad_report.json', ok=False)
    assert adapted_bytes == (target / 'test_a.py').read_bytes()
    runner = output / 'run_pytest.sh'
    run('generate_runner.py', '--test-root', target, '--output', runner)
    assert json.loads(runner.with_suffix('.files.json').read_text(encoding='utf-8'))['test_files'] == manifest['test_files']
    selected_runner = output / 'selected.sh'
    run('generate_runner.py', '--test-root', target, '--output', selected_runner, '--files', 'nested/test_a.py')
    assert json.loads(selected_runner.with_suffix('.files.json').read_text(encoding='utf-8'))['test_files'] == ['nested/test_a.py']
    bash = shutil.which('bash')
    if not bash and Path('C:/Program Files/Git/bin/bash.exe').exists():
        bash = 'C:/Program Files/Git/bin/bash.exe'
    if bash:
        subprocess.run([bash, '-n', str(runner)], check=True, capture_output=True)
        result = subprocess.run([bash, str(runner), str(target), '--dry-run'], capture_output=True, text=True)
        assert result.returncode == 0, result.stderr
        assert result.stdout.count('DRY RUN:') == len(manifest['test_files'])
        assert not (root / 'logs2').exists()
    logs = output / 'fixture_logs2'
    logs.mkdir()
    def log(name, file, body, rc=0):
        text = f'Test file: {file}\nCommand: python3 -m pytest -vv {file}\n{body}\n'
        if rc is not None:
            text += f'===== exit code: {rc} =====\n'
        path = logs / name
        path.write_text(text, encoding='utf-8')
        return path
    log('00001.log', 'test_a.py', '''test_a.py::TestBaseNPU::test_mix_npu_float32 PASSED [ 10%]
test_a.py::TestBaseNPU::test_mix_npu_float16 FAILED [ 20%]
test_a.py::TestBaseNPU::test_skip[one] PASSED [ 30%]
test_a.py::TestBaseNPU::test_skip[two] SKIPPED (requires feature) [ 40%]
test_a.py::TestBaseNPU::test_pass[with space] PASSED [ 50%]
test_a.py::Another::test_pass PASSED [ 60%]
test_a.py::TestBaseNPU::test_error ERROR [ 70%]
test_a.py::TestBaseNPU::test_xfail XFAIL (known bug) [ 80%]
test_a.py::TestBaseNPU::test_xpass XPASS [ 90%]
test_a.py::TestBaseNPU::test_addmm FAILED [100%]
==================== FAILURES ====================
________________ TestBaseNPU.test_mix_npu_float16 ________________
E   AssertionError: incorrect output
________________ TestBaseNPU.test_addmm ________________
E   RuntimeError: addmm failed
================ short test summary info ================
FAILED test_a.py::TestBaseNPU::test_mix_npu_float16 - AssertionError: incorrect output
FAILED test_a.py::TestBaseNPU::test_addmm - RuntimeError: addmm failed
ERROR test_a.py::TestBaseNPU::test_error - RuntimeError: setup failed''', 1)
    log('00002.log', 'nested/test_a.py', 'nested/test_a.py::test_pass FAILED [100%]\n=== short test summary info ===\nFAILED nested/test_a.py::test_pass - AssertionError: nested failure', 1)
    log('00003.log', 'test_collect.py', 'collected 0 items / 1 error\nE   ImportError: missing dependency\n=== short test summary info ===\nERROR test_collect.py - ImportError: missing dependency', 2)
    log('00004.log', 'test_interrupt.py', 'test_interrupt.py::test_pass PASSED [ 50%]', None)
    log('00005.log', 'test_exit.py', 'test_exit.py::test_pass PASSED [ 50%]', 2)
    results = output / 'pytest_results.xlsx'
    run('summarize_pytest_logs.py', '--input', logs, '--output', results, '--format', 'xlsx')
    raw = load_workbook(results, read_only=True)
    statuses = Counter(row[2] for row in raw.active.iter_rows(min_row=2, values_only=True))
    assert statuses['XFAIL'] == 1 and statuses['XPASS'] == 1 and statuses['ERROR'] >= 2
    raw.close()
    counts = fill(inventory, results, inventory)
    book = load_workbook(inventory)
    sheet = book.active
    assert all(sheet.cell(r, c).value == v for (r, c), v in base_cells.items())
    assert sheet['A2'].number_format == '@'
    assert book['preserved']['B2'].value == 'keep'
    found = {(row[0], row[1]): (row[3], row[4]) for row in sheet.iter_rows(min_row=2, values_only=True)}
    for name, expected in dict(test_mix='FAILED', test_skip='SKIPPED', test_pass='PASSED', test_missing='NOT_RUN',
                               test_error='ERROR', test_xfail='XFAIL', test_xpass='XPASS', test_add='FAILED').items():
        assert found[('test_a.py', name)][0] == expected, (name, found)
    assert found[('nested/test_a.py', 'test_pass')][0] == 'FAILED'
    assert 'incorrect output' in found[('test_a.py', 'test_mix')][1]
    assert 'missing dependency' in found[('test_collect.py', 'test_pass')][1]
    assert found[('test_interrupt.py', 'test_pass')][0] == 'NOT_RUN'
    assert found[('test_exit.py', 'test_pass')][0] == 'NOT_RUN'
    assert inventory.with_name('待测用例.before_fill.xlsx').exists()
    book.close()
    # Missing file results must not become another file's same-named method.
    assert found[('test_a.py', 'test_missing')][0] == 'NOT_RUN'
    # Check log variants (ANSI, split status, teardown ERROR after PASSED).
    variant = root / 'variant.log'
    variant.write_text('Test file: test_v.py\n\x1b[32mtest_v.py::TestV::test_one\x1b[0m\nwarning\nPASSED [100%]\n=== short test summary info ===\nERROR test_v.py::TestV::test_one - RuntimeError: teardown\n===== exit code: 1 =====\n', encoding='utf-8')
    parsed = parse_log(variant)
    assert len(parsed) == 1 and parsed[0].status == 'ERROR' and 'teardown' in parsed[0].error
    return dict(status='PASS', real_npu_tests_run=False, network_used=False,
                static_definitions=13, test_files=5, raw_statuses=dict(statuses),
                filled_statuses=counts, bash_syntax_and_dry_run='PASS' if bash else 'NOT_AVAILABLE')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path)
    args = parser.parse_args()
    if args.output_dir:
        if args.output_dir.exists():
            parser.error('Fixture output directory must not exist')
        result = check(args.output_dir.resolve())
        write_json(args.output_dir / 'fixture_validation.json', result)
    else:
        with tempfile.TemporaryDirectory(prefix='inductor_npu_fixture_') as temporary:
            result = check(Path(temporary))
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
