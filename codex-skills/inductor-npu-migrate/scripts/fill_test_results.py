#!/usr/bin/env python3
"""Fill only 执行结果 and 失败原因 using file-aware testcase-prefix aggregation."""
import argparse
from collections import Counter, defaultdict
from pathlib import Path
import shutil
from openpyxl import load_workbook

ALIASES = {
    'file': ('file_name', '文件名', '文件名称'),
    'case': ('test_case', '用例名', '测试用例', 'test_name'),
    'class': ('class_name', '类名'),
    'status': ('状态', 'status', '执行结果'),
    'reason': ('报错信息', 'error', '失败原因'),
    'node': ('node_id', 'nodeid'),
}


def columns(sheet):
    return {str(c.value).strip(): c.column for c in sheet[1] if c.value is not None}


def find_col(headers, field, required=True):
    for name in ALIASES[field]:
        if name in headers:
            return headers[name]
    if required:
        raise ValueError(f'Missing {field} column; accepted headers: {ALIASES[field]}')
    return None


def normalize_file(value):
    value = str(value or '').replace('\\', '/').removeprefix('./')
    return value.split('/test/inductor/')[-1].removeprefix('test/inductor/').removeprefix('inductor/')


def aggregate(matches):
    statuses = {m['status'] for m in matches}
    if 'FAILED' in statuses:
        selected = 'FAILED'
    elif 'SKIPPED' in statuses:
        selected = 'SKIPPED'
    elif statuses == {'PASSED'}:
        selected = 'PASSED'
    else:
        selected = next((s for s in ('ERROR', 'NOT_RUN', 'XFAIL', 'XPASS') if s in statuses), 'UNKNOWN')
    reasons = list(dict.fromkeys(m['reason'] for m in matches
                               if m['status'] == selected and m['reason']))
    return selected, '\n'.join(reasons)[:32767] if selected != 'PASSED' else ''


def fill(inventory, results, output, sheet_name=None, results_sheet=None, unmatched='NOT_RUN'):
    inventory, results, output = map(Path, (inventory, results, output))
    if output.resolve() == results.resolve():
        raise ValueError('Output must not overwrite the result workbook')
    workbook = load_workbook(inventory)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    raw = load_workbook(results, read_only=True, data_only=True)
    source = raw[results_sheet] if results_sheet else raw.active
    ih, rh = columns(sheet), columns(source)
    fc, cc = find_col(ih, 'file'), find_col(ih, 'case')
    rf, rc, rs = [find_col(rh, f) for f in ('file', 'case', 'status')]
    rr, rn = find_col(rh, 'reason', False), find_col(rh, 'node', False)
    index, file_issues = defaultdict(list), defaultdict(list)
    for row in source.iter_rows(min_row=2, values_only=True):
        file = normalize_file(row[rf - 1])
        name = str(row[rc - 1] or '')
        item = dict(name=name.rsplit('::', 1)[-1], status=str(row[rs - 1] or 'UNKNOWN').upper(),
                    reason=str(row[rr - 1] or '') if rr else '', node=str(row[rn - 1] or '') if rn else '')
        if '::' in item['node']:
            file = normalize_file(item['node'].split('::')[0])
        if (name.startswith('<') or (item['node'] and '::' not in item['node'])
                or name.endswith('.py') or not name.startswith('test_')):
            file_issues[file].append(item)
        else:
            index[file].append(item)
    # Retain all source columns and cells; add only the two requested columns if absent.
    for name in ('执行结果', '失败原因'):
        if name not in ih:
            ih[name] = sheet.max_column + 1
            sheet.cell(1, ih[name], name)
    counts = Counter()
    basename_index = defaultdict(set)
    for key in set(index) | set(file_issues):
        basename_index[Path(key).name].add(key)
    for row in range(2, sheet.max_row + 1):
        file = normalize_file(sheet.cell(row, fc).value)
        name = str(sheet.cell(row, cc).value or '').strip()
        if not name:
            continue
        key = file
        if key not in index and key not in file_issues:
            alternatives = basename_index[Path(file).name]
            if len(alternatives) == 1:
                key = next(iter(alternatives))
            elif len(alternatives) > 1:
                raise ValueError(f'Ambiguous result filename: {file}; keep relative paths/node_id')
        matches = [m for m in index[key] if m['name'].startswith(name)]
        # Default follows the requested prefix rule across generated device classes.
        # A declared generic class cannot safely be matched to guessed generated names.
        # Distinct same-name definitions in different classes are intentionally aggregated.
        if not matches:
            details = '; '.join(dict.fromkeys(m['reason'] for m in file_issues[key] if m['reason']))
            status, reason = unmatched, details or '未找到匹配实测结果（未运行、未收集或选择范围外）'
        else:
            status, reason = aggregate(matches)
            if status == 'PASSED' and file_issues[key]:
                status = 'NOT_RUN'
                reason = '; '.join(dict.fromkeys(m['reason'] for m in file_issues[key] if m['reason']))
        sheet.cell(row, ih['执行结果'], status)
        sheet.cell(row, ih['失败原因'], reason[:32767])
        counts[status or '(blank)'] += 1
    raw.close()
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.resolve() == inventory.resolve():
        backup = inventory.with_name(inventory.stem + '.before_fill.xlsx')
        if not backup.exists():
            shutil.copy2(inventory, backup)
    elif output.exists():
        raise ValueError(f'Refusing to overwrite existing output: {output}')
    temp = output.with_name(output.stem + '.writing.xlsx')
    if temp.exists():
        raise ValueError(f'Previous incomplete save exists: {temp}')
    workbook.save(temp)
    temp.replace(output)
    workbook.close()
    return dict(counts)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inventory', type=Path, default=Path('待测用例.xlsx'))
    parser.add_argument('--results', type=Path, required=True)
    parser.add_argument('--output', type=Path, help='Omit to update inventory in place, with backup')
    parser.add_argument('--sheet')
    parser.add_argument('--results-sheet')
    parser.add_argument('--unmatched', choices=('NOT_RUN', 'blank'), default='NOT_RUN')
    args = parser.parse_args()
    print(fill(args.inventory, args.results, args.output or args.inventory,
               args.sheet, args.results_sheet, '' if args.unmatched == 'blank' else 'NOT_RUN'))


if __name__ == '__main__':
    main()
