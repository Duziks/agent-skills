#!/usr/bin/env python3
"""Extract pytest case results from log files into one Excel workbook or CSV.

Usage:
    python3 summarize_pytest_logs.py
    python3 summarize_pytest_logs.py --input ./logs2 --output ./pytest_results.xlsx
    python3 summarize_pytest_logs.py --input ./test_custom_lowering.log
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None


RESULT_STATUSES = ("PASSED", "SKIPPED", "FAILED", "ERROR", "XFAIL", "XPASS")
STATUS_PATTERN = "|".join(RESULT_STATUSES)
NODEID_LINE = re.compile(r"^\s*(?P<node>\S+\.py(?:::[^\s]+)+)(?P<rest>.*)$")
STATUS_AT_START = re.compile(r"^\s*(?P<status>" + STATUS_PATTERN + r")\b")
STATUS_TRAILING = re.compile(
    r"\b(?P<status>" + STATUS_PATTERN + r")(?=\s*(?:\[|\(|$))"
)
COMMAND_LINE = re.compile(r"^Command:\s+pytest\b(?P<arguments>.*)$")
COLLECTED_ZERO = re.compile(r"^\s*collected\s+0\s+items\b", re.IGNORECASE)
NO_TESTS_RAN = re.compile(r"no tests ran", re.IGNORECASE)
EXIT_CODE = re.compile(r"^===== exit code:\s*(?P<code>\d+)\s*=====$")
# Pytest normally uses several underscores around a failure title, but some
# terminal/reporting configurations emit a single underscore on each side.
FAILURE_HEADER = re.compile(r"^\s*_+\s*(?P<name>.+?)\s*_+\s*$")
SECTION_HEADER = re.compile(
    r"^=+\s+(?:FAILURES|ERRORS|short test summary info)\s+=+$", re.IGNORECASE
)
SUMMARY_RESULT_LINE = re.compile(
    r"^\s*(?P<status>FAILED|ERROR)(?:\s+\[[^\]]+\])?\s+"
    r"(?P<node>.+?\.py(?:::.+?)?)(?:\s+-\s+(?P<reason>.*))?\s*$"
)
EXCEPTION_LINE = re.compile(
    r"^\s*(?:E\s+)?(?P<error>"
    r"(?:[A-Za-z_][\w.]*(?:Error|Exception|Failure|Interrupt|Exit|Unsupported)|"
    r"AssertionError|Failed)"
    r"(?::\s*.*)?"
    r")$"
)
ANSI_ESCAPE = re.compile(r"\x1B(?:[@-Z\\-_]|\[[0-?]*[ -/]*[@-~])")
CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


@dataclass
class CaseResult:
    file_name: str
    test_name: str
    status: str
    error: str = ""
    node_id: str = ""


@dataclass
class FailureBlock:
    name: str
    error: str = ""


def normalize_status(status: str) -> str:
    """Keep raw pytest outcomes; special outcomes are never silently called PASSED."""
    return status


def test_name_from_node(node: str) -> str:
    return node.rsplit("::", 1)[-1]


def failure_name_from_header(header: str) -> str | None:
    """Turn a pytest failure heading into the function name used in the result table."""
    header = re.sub(r"^ERROR\s+(?:at\s+setup\s+of|at\s+teardown\s+of)\s+", "", header)
    match = re.search(r"((?:[A-Za-z_][\w.]*\.)?test_[A-Za-z0-9_\[\]-]+)$", header)
    return match.group(1) if match else header.strip() or None


def parse_failure_blocks(lines: list[str]) -> list[FailureBlock]:
    """Return failure/error blocks and one concise exception line from each block."""
    blocks: list[FailureBlock] = []
    current_name: str | None = None
    current_lines: list[str] = []

    def save_current() -> None:
        if current_name is None:
            return

        candidates: list[str] = []
        fallback: list[str] = []
        for line in current_lines:
            clean_line = line.strip()
            match = EXCEPTION_LINE.match(clean_line)
            if match:
                candidates.append(match.group("error"))
            elif clean_line.startswith("E ") or clean_line.startswith("E\t"):
                fallback.append(re.sub(r"^E\s+", "", clean_line))
        error = candidates[0] if candidates else (fallback[-1] if fallback else "")
        blocks.append(FailureBlock(current_name, error))

    for line in lines:
        header = FAILURE_HEADER.match(line)
        if header:
            save_current()
            current_name = failure_name_from_header(header.group("name"))
            current_lines = []
        elif current_name is not None:
            current_lines.append(line)
    save_current()
    return blocks


def first_exception_line(lines: list[str]) -> str:
    for line in lines:
        match = EXCEPTION_LINE.match(line.strip())
        if match:
            return match.group("error")
    for line in lines:
        clean_line = line.strip()
        if clean_line.startswith("E ") or clean_line.startswith("E\t"):
            return re.sub(r"^E\s+", "", clean_line)
    return ""


def file_name_from_log(log_path: Path, lines: list[str]) -> str:
    for line in lines:
        if line.startswith("Test file: "):
            return line.removeprefix("Test file: ").strip()
        match = COMMAND_LINE.match(line)
        if match:
            arguments = re.findall(r"'[^']*'|\"[^\"]*\"|\S+", match.group("arguments"))
            for argument in reversed(arguments):
                argument = argument.strip("'\"")
                path_match = re.match(r"(?P<path>.+?\.py)(?:::.+)?$", argument)
                if path_match:
                    return path_match.group("path").replace("\\", "/").rsplit("/", 1)[-1]

    stem = log_path.stem
    return stem if stem.endswith(".py") else f"{stem}.py"


def parse_log(log_path: Path) -> list[CaseResult]:
    lines = [
        CONTROL_CHARS.sub("", ANSI_ESCAPE.sub("", line))
        for line in log_path.read_text(encoding="utf-8", errors="replace").splitlines()
    ]
    file_name = file_name_from_log(log_path, lines)
    failure_start = next(
        (index for index, line in enumerate(lines) if SECTION_HEADER.match(line)), len(lines)
    )
    execution_lines = lines[:failure_start]
    failure_blocks = parse_failure_blocks(lines[failure_start:])

    results: list[CaseResult] = []
    result_index_by_node: dict[str, int] = {}
    pending_node: str | None = None

    def add_result(node: str, status: str) -> None:
        result = CaseResult(
            file_name,
            test_name_from_node(node),
            normalize_status(status),
            node_id=node,
        )
        if node in result_index_by_node:
            index = result_index_by_node[node]
            previous = results[index]
            priority = {"FAILED": 7, "ERROR": 6, "SKIPPED": 5, "XFAIL": 4, "XPASS": 3, "PASSED": 1}
            if priority[result.status] >= priority[previous.status]:
                result.error = previous.error
                results[index] = result
        else:
            result_index_by_node[node] = len(results)
            results.append(result)

    for line in execution_lines:
        # Parametrized IDs may contain spaces, so parse complete verbose lines first.
        full_line = re.match(r"^\s*(?P<node>.+?\.py::.+?)\s+(?P<status>" + STATUS_PATTERN + r")\b(?P<rest>.*)$", line)
        if full_line:
            add_result(full_line.group("node"), full_line.group("status"))
            if full_line.group("status") in {"SKIPPED", "XFAIL", "XPASS"}:
                detail = re.sub(r"\s*\[\s*\d+%\]\s*$", "", full_line.group("rest")).strip(' ()')
                results[result_index_by_node[full_line.group("node")]].error = detail
            pending_node = None
            continue
        node_match = NODEID_LINE.match(line)
        if node_match:
            node = node_match.group("node")
            remainder = node_match.group("rest")
            # A warning can be printed between the nodeid and its final status.
            status_match = STATUS_TRAILING.search(remainder)
            if status_match:
                add_result(node, status_match.group("status"))
            else:
                pending_node = node
            if status_match:
                pending_node = None
            continue

        status_match = STATUS_AT_START.match(line)
        if status_match and pending_node is not None:
            add_result(pending_node, status_match.group("status"))
            pending_node = None

    for line in lines[failure_start:]:
        summary_match = SUMMARY_RESULT_LINE.match(line)
        if summary_match:
            add_result(summary_match.group("node"), summary_match.group("status"))
            reason = summary_match.group("reason")
            if reason:
                results[result_index_by_node[summary_match.group("node")]].error = reason

    failed_results = [result for result in results if result.status in {"FAILED", "ERROR"}]
    unused_blocks = [block for block in failure_blocks if block.error]
    used_blocks: set[int] = set()

    def block_matches(result: CaseResult, block: FailureBlock) -> bool:
        if not result.node_id or not block.name:
            return False
        result_tail = result.node_id.split(".py::", 1)[-1].replace("::", ".")
        block_name = block.name.replace("::", ".")
        return (
            result_tail == block_name
            or result_tail.endswith(f".{block_name}")
            or block_name.endswith(f".{result_tail}")
        )

    for result in failed_results:
        for block_index, block in enumerate(unused_blocks):
            if block_index not in used_blocks and block_matches(result, block):
                result.error = block.error
                used_blocks.add(block_index)
                break

    for result in results:
        if result.status in {"FAILED", "ERROR"}:
            if not result.error:
                result.error = "Failure details not found in log."

    collected_zero = any(COLLECTED_ZERO.match(line) for line in lines)
    no_tests_ran = any(NO_TESTS_RAN.search(line) for line in lines)
    if (collected_zero or no_tests_ran) and not results:
        exit_code = next(
            (match.group("code") for line in lines if (match := EXIT_CODE.match(line))),
            None,
        )
        detail = "collected 0 items" if collected_zero else "no tests ran"
        if exit_code is not None:
            detail += f" (pytest exit code: {exit_code})"
        collection_error = first_exception_line(lines)
        if collection_error:
            detail += f"; {collection_error}"
        results.append(CaseResult(file_name, "<未收集到用例>", "ERROR", detail))

    if not results:
        results.append(
            CaseResult(
                file_name,
                "<日志解析失败>",
                "ERROR",
                "No PASSED, SKIPPED, or FAILED test result was found in this log.",
            )
        )
    if results and not any(EXIT_CODE.match(line) for line in lines):
        results.append(CaseResult(file_name, "<运行未完整结束>", "NOT_RUN",
                                  "Runner exit marker missing; log may be interrupted or truncated."))
    exit_codes = [int(m.group('code')) for line in lines if (m := EXIT_CODE.match(line))]
    if exit_codes and exit_codes[-1] not in {0, 1} and not any(r.test_name.startswith('<') for r in results):
        results.append(CaseResult(file_name, "<文件执行异常>", "ERROR",
                                  f'pytest exit code {exit_codes[-1]}; file did not complete normally'))
    return results


def discover_logs(input_path: Path) -> list[Path]:
    if input_path.is_file():
        return [input_path]
    if not input_path.is_dir():
        raise FileNotFoundError(f"Input path does not exist: {input_path}")
    return sorted(path for path in input_path.rglob("*.log") if path.is_file())


HEADERS = ["文件名", "用例名", "状态", "报错信息", "node_id"]


def write_csv(results: list[CaseResult], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(HEADERS)
        writer.writerows(
            [result.file_name, result.test_name, result.status, result.error, result.node_id]
            for result in results
        )


def write_xlsx(results: list[CaseResult], output_path: Path) -> None:
    if openpyxl is None:
        raise RuntimeError(
            "openpyxl is not installed; use --format csv or install it with "
            "python3 -m pip install openpyxl"
        )

    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Pytest Results"
    sheet.append(HEADERS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for result in results:
        sheet.append([result.file_name, result.test_name, result.status, result.error, result.node_id])

    for row in sheet.iter_rows(min_row=2):
        row[2].alignment = Alignment(horizontal="center")
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
        if row[2].value in {"FAILED", "ERROR"}:
            row[2].fill = PatternFill("solid", fgColor="F4CCCC")
        elif row[2].value == "SKIPPED":
            row[2].fill = PatternFill("solid", fgColor="FFF2CC")
        elif row[2].value == "PASSED":
            row[2].fill = PatternFill("solid", fgColor="D9EAD3")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 55
    sheet.column_dimensions["C"].width = 14
    sheet.column_dimensions["D"].width = 120
    sheet.column_dimensions["E"].width = 100

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_results(
    results: list[CaseResult], output_path: Path, output_format: str
) -> Path:
    """Write xlsx when available, otherwise a directly-openable UTF-8 CSV."""
    output_format = output_format.lower()
    if output_format not in {"auto", "xlsx", "csv"}:
        raise ValueError(f"Unsupported output format: {output_format}")

    if output_format == "auto":
        output_format = "csv" if output_path.suffix.lower() == ".csv" else "xlsx"
        if output_format == "xlsx" and openpyxl is None:
            output_path = output_path.with_suffix(".csv")
            output_format = "csv"
            print(
                "Warning: openpyxl is not installed; writing CSV instead: "
                f"{output_path}",
                file=sys.stderr,
            )
    elif output_format == "csv" and output_path.suffix.lower() != ".csv":
        output_path = output_path.with_suffix(".csv")
    elif output_format == "xlsx" and output_path.suffix.lower() != ".xlsx":
        output_path = output_path.with_suffix(".xlsx")

    if output_format == "csv":
        write_csv(results, output_path)
    else:
        write_xlsx(results, output_path)
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract pytest results from log files into one Excel workbook or CSV."
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=Path("logs2"),
        help="A .log file or directory of .log files (default: ./logs2)",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("pytest_results.xlsx"),
        help="Output Excel workbook path (default: ./pytest_results.xlsx)",
    )
    parser.add_argument(
        "--format",
        choices=("auto", "xlsx", "csv"),
        default="auto",
        help="Output format (default: auto; falls back to CSV without openpyxl)",
    )
    args = parser.parse_args()

    try:
        log_files = discover_logs(args.input)
    except FileNotFoundError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2

    if not log_files:
        print(f"Error: no .log files found in {args.input}", file=sys.stderr)
        return 2

    results: list[CaseResult] = []
    for log_file in log_files:
        results.extend(parse_log(log_file))

    try:
        actual_output = write_results(results, args.output, args.format)
    except (RuntimeError, OSError, ValueError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2
    print(
        f"Processed {len(log_files)} log file(s), extracted {len(results)} row(s): "
        f"{actual_output}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
